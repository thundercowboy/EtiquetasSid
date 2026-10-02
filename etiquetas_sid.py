"""Etiquetas Sid: dados locais, edição individual e etiquetas A4 em PDF."""
import json
import os
import re
import sys
from datetime import date, datetime
from pathlib import Path
import tkinter as tk
from tkinter import ttk, messagebox, filedialog

BASE = Path(sys.executable).parent if getattr(sys, 'frozen', False) else Path(__file__).resolve().parent
DEFAULT_CONFIG = dict(margem_esquerda_mm=4.65, margem_superior_mm=13.1,
                      espaco_horizontal_mm=2.5, espaco_vertical_mm=0.0)


def normalizar_numero(value):
    if value is None or isinstance(value, bool):
        raise ValueError('Número do paciente vazio ou inválido.')
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if not re.fullmatch(r'[0-9]+', text):
        raise ValueError('O número do paciente deve conter somente dígitos.')
    return text.lstrip('0') or '0'


def carregar_pacientes(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.worksheets[0].iter_rows(max_col=16, values_only=True)
        headers = [str(v).strip().upper() if v is not None else '' for v in next(rows, ())]
        required = ['NÚMERO VIDAS', 'NOME COMPLETO', 'DATA NASCIMENTO']
        for name in required:
            if headers.count(name) != 1:
                raise ValueError(f'O cabeçalho {name} deve aparecer uma vez na primeira linha (A:P).')
        indices = [headers.index(name) for name in required]
        patients = {}
        for row_number, row in enumerate(rows, 2):
            number, name, birth = [row[i] for i in indices]
            if number is None or str(number).strip() == '':
                continue
            try:
                key = normalizar_numero(number)
            except ValueError as exc:
                raise ValueError(f'Linha {row_number}: {exc}') from exc
            if key in patients:
                raise ValueError(f'Número de paciente duplicado na linha {row_number}.')
            if name is None or not str(name).strip():
                raise ValueError(f'Nome ausente na linha {row_number}.')
            if isinstance(birth, (datetime, date)):
                birth = birth.strftime('%d/%m/%Y')
            elif birth is not None:
                try:
                    birth = datetime.strptime(str(birth).strip(), '%d/%m/%Y').strftime('%d/%m/%Y')
                except ValueError as exc:
                    raise ValueError(f'Data de nascimento inválida na linha {row_number}. Use uma data do Excel ou dd/mm/aaaa.') from exc
            else:
                raise ValueError(f'Data de nascimento ausente na linha {row_number}.')
            patients[key] = dict(numero=key, nome=str(name).strip(), nascimento=birth)
        return patients
    finally:
        workbook.close()


def carregar_responsaveis():
    path = BASE / 'responsaveis.txt'
    if not path.exists():
        path.write_text('Roberta Rezende - Nucleo 1\nAmanda Altino - Nucleo 2\nLuana Lopes - Nucleo 3\n', encoding='utf-8')
    return list(dict.fromkeys(line.strip() for line in path.read_text(encoding='utf-8-sig').splitlines() if line.strip()))


def gerar_pdf(path, labels, config, start=1):
    from reportlab.pdfgen import canvas
    from reportlab.lib.units import mm
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase.pdfmetrics import stringWidth
    left, top, gap_x, gap_y = [float(config[k]) for k in DEFAULT_CONFIG]
    if min(left, top, gap_x, gap_y) < 0 or left + 198.2 + gap_x > 210 or top + 270.8 + 3 * gap_y > 297:
        raise ValueError('As margens e os espaçamentos ultrapassam a folha A4.')
    if not 1 <= start <= 8:
        raise ValueError('A posição inicial deve estar entre 1 e 8.')
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle('Etiquetas Sid')
    def fitted(text, x, y, width, font='Helvetica', size=10):
        while stringWidth(text, font, size) > width and size > 6:
            size -= .25
        if stringWidth(text, font, size) > width:
            raise ValueError('Um dos textos é longo demais para a etiqueta. Reduza o conteúdo.')
        c.setFont(font, size)
        c.setFillColorRGB(0, 0, 0)
        c.drawString(x, y, text)
    for index, label in enumerate(labels, start - 1):
        if index and index % 8 == 0:
            c.showPage()
        slot = index % 8
        x = (left + (slot % 2) * (99.1 + gap_x) + 3) * mm
        y = A4[1] - (top + (slot // 2) * (67.7 + gap_y) + 6) * mm
        width = 93.1 * mm
        fitted(label['nome'].upper(), x, y, width, 'Helvetica-Bold', 11)
        def field(title, value, dx, dy, w):
            fitted(title, x + dx * mm, y - dy * mm, w * mm, 'Helvetica-Bold', size=6.5)
            fitted(value, x + dx * mm, y - (dy + 4) * mm, w * mm, 'Courier-Bold', 9)
        field('Nº MAIS VIDAS', label['numero'], 0, 7, 45)
        field('DATA NASC.', label['nascimento'], 48, 7, 45)
        field('PERÍODO REF.', label['periodo'], 0, 19, 93.1)
        field('DATA RECOLHIMENTO', '', 0, 31, 45)
        field('RESP. ENVIO', label['responsavel'], 48, 31, 45)
        field('RESP. RECOLHIMENTO', '', 0, 43, 93.1)
        c.setStrokeColorRGB(.25, .25, .25)
        c.setLineWidth(.5)
        for offset in (2, 14, 26, 38, 50):
            c.line(x, y - offset * mm, x + width, y - offset * mm)
    c.save()


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title('Etiquetas Sid')
        assets = Path(getattr(sys, '_MEIPASS', BASE))
        icon_path = assets / 'imgs' / 'icon.png'
        if icon_path.exists():
            self.app_icon = tk.PhotoImage(file=str(icon_path))
            self.iconphoto(True, self.app_icon)
        self.geometry('980x650')
        self.minsize(850, 580)
        self.style = ttk.Style(self)
        self.style.theme_use('clam')
        self.dark_mode = tk.BooleanVar(value=True)
        self.apply_theme()
        self.labels = []
        self.config_values = DEFAULT_CONFIG.copy()
        try:
            if (BASE / 'config.json').exists():
                self.config_values.update(json.loads((BASE / 'config.json').read_text(encoding='utf-8')))
        except (ValueError, OSError):
            messagebox.showwarning('Configuração', 'Configuração inválida. Usando medidas iniciais.')
        frame = ttk.Frame(self, padding=16)
        frame.pack(fill='both', expand=True)
        header = ttk.Frame(frame)
        header.pack(fill='x')
        ttk.Label(header, text='Etiquetas Sid', font=('Segoe UI', 20, 'bold')).pack(side='left')
        ttk.Checkbutton(header, text='Modo escuro', variable=self.dark_mode, command=self.apply_theme, style='Toggle.TCheckbutton').pack(side='right')
        ttk.Label(frame, text='BD/BD.xlsx • primeira aba • 8 etiquetas por folha A4').pack(anchor='w', pady=(0, 12))
        search = ttk.Frame(frame)
        search.pack(fill='x')
        ttk.Label(search, text='Números dos pacientes:').pack(side='left')
        self.numbers = ttk.Entry(search)
        self.numbers.pack(side='left', fill='x', expand=True, padx=8)
        self.numbers.bind('<Return>', lambda _: self.search())
        ttk.Button(search, text='Adicionar pacientes', command=self.search).pack(side='left')
        ttk.Label(frame, text='Separe os números por espaço, vírgula ou ponto e vírgula. Zeros à esquerda são aceitos.').pack(anchor='w', pady=5)
        self.tree = ttk.Treeview(frame, columns=('numero', 'nome', 'nascimento', 'periodo', 'responsavel'), show='headings', height=10, selectmode='extended')
        for key, title, width in [('numero','Nº MAIS VIDAS',110), ('nome','Nome completo',270), ('nascimento','Nascimento',100), ('periodo','Período ref.',160), ('responsavel','Resp. envio',160)]:
            self.tree.heading(key, text=title)
            self.tree.column(key, width=width)
        self.tree.pack(fill='both', expand=True, pady=8)
        self.tree.bind('<<TreeviewSelect>>', self.select)
        ttk.Label(frame, text='Use Ctrl ou Shift para selecionar várias etiquetas. Campos vazios preservam os valores na edição em grupo.').pack(anchor='w')
        ttk.Button(frame, text='Selecionar todas', command=lambda: self.tree.selection_set(self.tree.get_children())).pack(anchor='w')
        editor = ttk.LabelFrame(frame, text='Preenchimento das etiquetas selecionadas', padding=10)
        editor.pack(fill='x')
        ttk.Label(editor, text='PERÍODO REF.:').grid(row=0, column=0, sticky='w')
        self.period = ttk.Entry(editor, width=40)
        self.period.grid(row=0, column=1, padx=8)
        ttk.Label(editor, text='RESP. ENVIO:').grid(row=0, column=2)
        self.responsible = ttk.Combobox(editor, state='readonly', width=22)
        self.responsible.grid(row=0, column=3, padx=8)
        ttk.Button(editor, text='Aplicar às selecionadas', command=self.save_fields).grid(row=1, column=1, pady=8, sticky='w')
        ttk.Button(editor, text='Remover selecionadas', command=self.remove).grid(row=1, column=3)
        controls = ttk.Frame(frame)
        controls.pack(fill='x', pady=12)
        ttk.Button(controls, text='Recarregar responsáveis', command=self.reload).pack(side='left')
        ttk.Button(controls, text='Ajustar alinhamento', command=self.settings).pack(side='left', padx=8)
        ttk.Label(controls, text='Posição inicial:').pack(side='left')
        self.start = ttk.Spinbox(controls, from_=1, to=8, width=3)
        self.start.set('1')
        self.start.pack(side='left', padx=5)
        ttk.Button(controls, text='Gerar PDF / imprimir', command=self.export).pack(side='right')
        ttk.Label(frame, text='Campos de recolhimento ficam em branco. No leitor de PDF, imprima em tamanho real (100%).').pack(anchor='w')
        self.reload()

    def apply_theme(self):
        dark = self.dark_mode.get()
        bg, panel, text, border, active = (
            ('#171b23', '#252c38', '#f1f5f9', '#536174', '#364358') if dark else
            ('#f3f5f8', '#ffffff', '#172033', '#9aa7b8', '#e1e8f0'))
        self.configure(background=bg)
        self.style.configure('.', background=bg, foreground=text, font=('Segoe UI', 9))
        for name in ('TFrame', 'TLabel', 'TLabelframe', 'TLabelframe.Label', 'TCheckbutton', 'Toggle.TCheckbutton'):
            self.style.configure(name, background=bg, foreground=text)
        self.style.configure('TLabelframe', bordercolor=border)
        self.style.configure('TButton', background=panel, foreground=text, bordercolor=border, padding=6)
        self.style.map('TButton', background=[('active', active)], foreground=[('disabled', '#8993a3'), ('active', text)])
        self.style.map('TCheckbutton', background=[('active', bg)], foreground=[('active', text)])
        self.style.map('Toggle.TCheckbutton', background=[('active', bg)], foreground=[('active', text)])
        for name in ('TEntry', 'TCombobox', 'TSpinbox'):
            self.style.configure(name, fieldbackground=panel, background=panel, foreground=text,
                                 insertcolor=text, bordercolor=border, arrowcolor=text,
                                 selectbackground='#2563eb', selectforeground='#ffffff')
            self.style.map(name, fieldbackground=[('readonly', panel), ('disabled', bg)],
                           foreground=[('readonly', text)], background=[('active', active)])
        self.style.configure('Treeview', background=panel, fieldbackground=panel, foreground=text,
                             bordercolor=border, rowheight=26)
        self.style.map('Treeview', background=[('selected', '#2563eb')], foreground=[('selected', '#ffffff')])
        self.style.configure('Treeview.Heading', background=active, foreground=text, bordercolor=border)
        self.style.map('Treeview.Heading', background=[('active', panel)])
        self.option_add('*TCombobox*Listbox.background', panel)
        self.option_add('*TCombobox*Listbox.foreground', text)
        self.option_add('*TCombobox*Listbox.selectBackground', '#2563eb')
        self.option_add('*TCombobox*Listbox.selectForeground', '#ffffff')
        # Update an already-created dropdown when switching themes.
        if hasattr(self, 'responsible'):
            popup = self.tk.call('ttk::combobox::PopdownWindow', self.responsible)
            self.tk.call(f'{popup}.f.l', 'configure', '-background', panel, '-foreground', text)
        for child in self.winfo_children():
            if isinstance(child, tk.Toplevel):
                child.configure(background=bg)

    def reload(self):
        try:
            self.responsible['values'] = carregar_responsaveis()
        except Exception as exc:
            messagebox.showerror('Responsáveis', str(exc))

    def refresh(self):
        self.tree.delete(*self.tree.get_children())
        for i, label in enumerate(self.labels):
            self.tree.insert('', 'end', iid=str(i), values=[label[k] for k in ('numero','nome','nascimento','periodo','responsavel')])

    def search(self):
        try:
            tokens = re.split(r'[\s,;]+', self.numbers.get().strip())
            if not tokens or not tokens[0]:
                raise ValueError('Informe pelo menos um número de paciente.')
            keys = list(dict.fromkeys(normalizar_numero(t) for t in tokens))
            patients = carregar_pacientes(BASE / 'BD' / 'BD.xlsx')
            missing = [k for k in keys if k not in patients]
            for key in keys:
                if key in patients:
                    self.labels.append(dict(patients[key], periodo='', responsavel=''))
            self.refresh()
            if missing:
                messagebox.showwarning('Busca', 'Pacientes não encontrados: ' + ', '.join(missing))
            self.numbers.delete(0, 'end')
        except Exception as exc:
            messagebox.showerror('Leitura da planilha', str(exc))

    def select(self, _=None):
        selection = self.tree.selection()
        if selection:
            labels = [self.labels[int(item)] for item in selection]
            periods = {label['periodo'] for label in labels}
            responsibles = {label['responsavel'] for label in labels}
            self.period.delete(0, 'end')
            self.period.insert(0, next(iter(periods)) if len(periods) == 1 else '')
            self.responsible.set(next(iter(responsibles)) if len(responsibles) == 1 else '')

    def save_fields(self):
        selection = self.tree.selection()
        if not selection:
            messagebox.showinfo('Etiqueta', 'Selecione uma etiqueta na lista.')
            return
        period = self.period.get().strip()
        responsible = self.responsible.get()
        for item in selection:
            index = int(item)
            if period or len(selection) == 1:
                self.labels[index]['periodo'] = period
            if responsible or len(selection) == 1:
                self.labels[index]['responsavel'] = responsible
            self.tree.item(item, values=[self.labels[index][k] for k in ('numero','nome','nascimento','periodo','responsavel')])

    def remove(self):
        if self.tree.selection():
            for index in sorted((int(item) for item in self.tree.selection()), reverse=True):
                del self.labels[index]
            self.refresh()

    def settings(self):
        dialog = tk.Toplevel(self)
        dialog.configure(background=self.cget('background'))
        dialog.title('Alinhamento em milímetros')
        entries = {}
        titles = ['Margem esquerda', 'Margem superior', 'Espaço entre colunas', 'Espaço entre linhas']
        for row, (key, title) in enumerate(zip(DEFAULT_CONFIG, titles)):
            ttk.Label(dialog, text=title, padding=8).grid(row=row, column=0)
            entry = ttk.Entry(dialog)
            entry.insert(0, str(self.config_values[key]))
            entry.grid(row=row, column=1, padx=8)
            entries[key] = entry
        def save():
            try:
                values = {k: float(e.get().replace(',', '.')) for k, e in entries.items()}
                l, t, x, y = values.values()
                if min(values.values()) < 0 or l + 198.2 + x > 210 or t + 270.8 + 3*y > 297:
                    raise ValueError('Medidas ultrapassam a folha A4.')
                (BASE / 'config.json').write_text(json.dumps(values, indent=2), encoding='utf-8')
                self.config_values = values
                dialog.destroy()
            except Exception as exc:
                messagebox.showerror('Alinhamento', str(exc), parent=dialog)
        ttk.Button(dialog, text='Salvar', command=save).grid(row=4, column=1, pady=10)

    def export(self):
        try:
            if self.tree.selection():
                self.save_fields()
            if not self.labels:
                raise ValueError('Adicione pelo menos uma etiqueta.')
            names = carregar_responsaveis()
            for index, label in enumerate(self.labels, 1):
                if not label['periodo'] or label['responsavel'] not in names:
                    raise ValueError(f'Preencha o período e selecione um responsável válido na etiqueta {index}.')
            start = int(self.start.get())
            path = filedialog.asksaveasfilename(defaultextension='.pdf', initialfile='etiquetas.pdf', filetypes=[('PDF','*.pdf')])
            if path:
                gerar_pdf(path, self.labels, self.config_values, start)
                messagebox.showinfo('PDF gerado', 'PDF gerado. Imprima em A4, tamanho real (100%), sem ajustar à página.')
                os.startfile(path)
        except Exception as exc:
            messagebox.showerror('Gerar etiquetas', str(exc))


if __name__ == '__main__':
    App().mainloop()
